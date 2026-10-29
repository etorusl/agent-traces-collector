from __future__ import annotations

import contextlib
import io
import json
import os
import shutil
import traceback
from pathlib import Path
from typing import Any

import httpx


class CappedOutput(io.TextIOBase):
    def __init__(self, limit: int = 200000):
        self.parts: list[str] = []
        self.limit, self.length, self.total = limit, 0, 0

    def write(self, text: str) -> int:
        self.total += len(text)
        kept = text[: max(0, self.limit - self.length)]
        self.parts.append(kept)
        self.length += len(kept)
        return len(text)

    def flush(self) -> None:
        pass

    def getvalue(self) -> str:
        result = "".join(self.parts)
        if self.total > self.length:
            result += f"\n[tool output truncated: {self.total} characters total]"
        return result


class GaiaTools:
    """A local text-oriented agent environment. Python is NOT a security sandbox."""

    def __init__(self, task: dict[str, Any], config: dict[str, Any], workspace: str):
        self.root = Path(workspace).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        os.chdir(self.root)
        self.config = config
        self.attachments = []
        for i, attachment in enumerate(task["attachments"]):
            source = Path(attachment)
            destination = self.root / "attachments" / f"{i}_{source.name}"
            destination.parent.mkdir(exist_ok=True)
            shutil.copy2(source, destination)
            self.attachments.append(str(destination.relative_to(self.root)))
        self.namespace: dict[str, Any] = {"__name__": "__agent__", "attachments": self.attachments}

    def setup(self) -> dict[str, Any]:
        search = self.config["benchmark"]["search_provider"]
        return {
            "attachments": self.attachments,
            "search_provider": search,
            "capabilities": "Python, text/PDF/DOCX/XLSX reading, HTTP fetch; text-only model",
        }

    def path(self, value: str) -> Path:
        path = (self.root / value).resolve()
        if not path.is_relative_to(self.root):
            raise ValueError("File tools accept only paths inside the task workspace")
        return path

    def execute(self, name: str, args: dict[str, Any]) -> str:
        if name == "python":
            output = CappedOutput()
            with contextlib.redirect_stdout(output), contextlib.redirect_stderr(output):
                try:
                    exec(compile(args["code"], "<agent>", "exec"), self.namespace)
                except BaseException:
                    traceback.print_exc(limit=8)
            return output.getvalue() or "[no output; use print() to inspect values]"
        if name == "read_file":
            return self.read_file(**args)
        if name == "fetch_url":
            return self.fetch_url(**args)
        if name == "search_web" and self.config["benchmark"]["search_provider"] == "brave":
            return self.search_web(**args)
        raise ValueError(f"Unknown or disabled tool: {name}")

    def read_file(self, path: str, offset: int = 0, limit: int = 20000) -> str:
        if offset < 0 or not 1 <= limit <= 100000:
            raise ValueError("offset >= 0 and 1 <= limit <= 100000 required")
        p = self.path(path)
        if p.stat().st_size > 30 * 1024 * 1024:
            raise ValueError("File exceeds 30 MiB; inspect it with Python in chunks")
        suffix = p.suffix.lower()
        if suffix == ".pdf":
            from pypdf import PdfReader

            text = "\n".join(
                f"[page {i + 1}]\n{page.extract_text() or ''}"
                for i, page in enumerate(PdfReader(p).pages)
            )
        elif suffix == ".docx":
            from docx import Document

            document = Document(p)
            text = "\n".join(
                [x.text for x in document.paragraphs]
                + [
                    "\t".join(c.text for c in row.cells)
                    for table in document.tables
                    for row in table.rows
                ]
            )
        elif suffix == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(p, read_only=True, data_only=True)
            try:
                parts = []
                for sheet in wb.worksheets:
                    parts.append(f"[sheet: {sheet.title}; first 2000 rows, 100 columns]")
                    parts.extend(
                        str(row)
                        for row in sheet.iter_rows(
                            max_row=2000, max_col=min(sheet.max_column, 100), values_only=True
                        )
                    )
                text = "\n".join(parts)
            finally:
                wb.close()
        elif suffix in {".png", ".jpg", ".jpeg", ".gif", ".mp3", ".mp4", ".wav", ".avi"}:
            raise ValueError("This configuration has no vision/audio tools; inspect using Python")
        elif suffix == ".zip":
            import zipfile

            with zipfile.ZipFile(p) as archive:
                text = "\n".join(archive.namelist())
        else:
            text = p.read_text(encoding="utf-8", errors="replace")
        return (
            f"[characters {offset}:{offset + limit}; total {len(text)}]\n"
            + text[offset : offset + limit]
        )

    def fetch_url(self, url: str, offset: int = 0, limit: int = 20000) -> str:
        if not url.startswith(("https://", "http://")):
            raise ValueError("Only HTTP(S) URLs are supported")
        if offset < 0 or not 1 <= limit <= 100000:
            raise ValueError("Invalid offset/limit")
        with httpx.Client(follow_redirects=True, timeout=30) as client:
            with client.stream(
                "GET", url, headers={"User-Agent": "agent-traces-collector/0.1"}
            ) as r:
                r.raise_for_status()
                chunks = bytearray()
                for chunk in r.iter_bytes():
                    chunks.extend(chunk)
                    if len(chunks) > 8 * 1024 * 1024:
                        raise ValueError("HTTP content exceeds 8 MiB")
                content_type = r.headers.get("content-type", "")
                final_url = str(r.url)
        if "pdf" in content_type:
            from pypdf import PdfReader

            text = "\n".join(
                page.extract_text() or "" for page in PdfReader(io.BytesIO(chunks)).pages
            )
        else:
            text = chunks.decode("utf-8", errors="replace")
            if "html" in content_type:
                from bs4 import BeautifulSoup

                soup = BeautifulSoup(text, "html.parser")
                for tag in soup(["script", "style", "nav"]):
                    tag.decompose()
                links = [
                    {"text": a.get_text(" ", strip=True), "href": a.get("href")}
                    for a in soup.find_all("a", href=True)
                ][:100]
                text = soup.get_text("\n", strip=True) + "\n[links]\n" + json.dumps(links)
        return (
            f"URL: {final_url}\n[total characters: {len(text)}]\n" + text[offset : offset + limit]
        )

    def search_web(self, query: str, count: int = 5) -> str:
        key = os.environ.get(self.config["benchmark"]["search_api_key_env"], "")
        if not key:
            raise ValueError("Brave search key is missing")
        response = httpx.get(
            "https://api.search.brave.com/res/v1/web/search",
            params={"q": query, "count": max(1, min(count, 20))},
            headers={"X-Subscription-Token": key},
            timeout=30,
        )
        response.raise_for_status()
        return json.dumps(response.json().get("web", {}).get("results", []), ensure_ascii=False)

    def close(self) -> None:
        pass
